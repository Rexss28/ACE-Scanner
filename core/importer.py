"""CSV/Excel importer for employee lists.

HR provides a file with employee names. This importer:
  - Reads the file
  - Cleans up the data
  - Auto-generates barcodes
  - Inserts into the database

Returns an ImportResult with a summary.
"""

import csv
from pathlib import Path

from core.data_access import DataAccess
from core.models import ImportResult


class EmployeeImporter:
    def __init__(self, db: DataAccess):
        self.db = db

    # ------------------------------------------------------------------
    # Public API
    # ------------------------------------------------------------------
    def import_from_csv(self, path: str, name_column: str = "full_name") -> ImportResult:
        """Import employees from a CSV file."""
        rows = self._read_csv(path, name_column)
        return self._import_rows(rows)

    def import_from_excel(self, path: str, name_column: str = "full_name") -> ImportResult:
        """Import employees from an Excel (.xlsx) file."""
        rows = self._read_excel(path, name_column)
        return self._import_rows(rows)

    # ------------------------------------------------------------------
    # File readers
    # ------------------------------------------------------------------
    def _read_csv(self, path: str, name_column: str) -> list[str]:
        """Read names from a CSV. Returns a list of cleaned names."""
        names = []
        with open(path, newline="", encoding="utf-8-sig") as f:
            reader = csv.DictReader(f)

            if name_column not in (reader.fieldnames or []):
                raise ValueError(
                    f"CSV is missing the '{name_column}' column. "
                    f"Found columns: {reader.fieldnames}"
                )

            for row in reader:
                name = (row.get(name_column) or "").strip()
                if name:
                    names.append(name)

        return names

    def _read_excel(self, path: str, name_column: str) -> list[str]:
        """Read names from an Excel file. Returns a list of cleaned names."""
        try:
            from openpyxl import load_workbook
        except ImportError:
            raise ImportError(
                "Excel support requires 'openpyxl'. Install it with:\n"
                "    pip install openpyxl"
            )

        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb.active

        rows_iter = ws.iter_rows(values_only=True)
        header = next(rows_iter, None)

        if header is None:
            raise ValueError("Excel file is empty.")

        # Find the column index matching name_column (case-insensitive)
        header_clean = [str(h).strip() if h else "" for h in header]
        try:
            col_idx = [h.lower() for h in header_clean].index(name_column.lower())
        except ValueError:
            raise ValueError(
                f"Excel file is missing the '{name_column}' column. "
                f"Found columns: {header_clean}"
            )

        names = []
        for row in rows_iter:
            if row is None or col_idx >= len(row):
                continue
            value = row[col_idx]
            if value is None:
                continue
            name = str(value).strip()
            if name:
                names.append(name)

        wb.close()
        return names

    # ------------------------------------------------------------------
    # Core import logic
    # ------------------------------------------------------------------
    def _import_rows(self, names: list[str]) -> ImportResult:
        """Insert each name into the DB with a generated barcode."""
        imported = 0
        skipped = 0
        errors: list[str] = []

        for name in names:
            try:
                barcode = self.db.generate_barcode()
                self.db.add_employee(barcode, name)
                imported += 1
            except Exception as e:
                skipped += 1
                errors.append(f"{name}: {e}")

        return ImportResult(
            total=len(names),
            imported=imported,
            skipped=skipped,
            errors=errors,
        )